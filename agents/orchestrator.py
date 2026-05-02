"""Orchestrator — coordinates search → tailor → submit pipeline."""
import datetime
import os
import re
import anyio
from pathlib import Path
from zoneinfo import ZoneInfo
from claude_agent_sdk import (
    query, ClaudeAgentOptions, ResultMessage, AssistantMessage, TextBlock,
    ThinkingBlock, ToolUseBlock, TaskStartedMessage, TaskProgressMessage,
)

_CYAN    = "\033[36m"
_GREEN   = "\033[32m"
_RED     = "\033[31m"
_MAGENTA = "\033[35m"
_DIM     = "\033[2m"
_RESET   = "\033[0m"

_ANSI_RE = re.compile(r'\033\[[0-9;]*m')  # strip existing ANSI codes before recoloring

_ICONS = {
    "[Orchestrator]": "🎬",
    "[Tailor]":       "✂️ ",
    "[Search]":       "🔍",
    "[Submitter]":    "📨",
}

_ERASE_LINE = "\r\033[2K"   # go to line start, erase entire line

_line_buf           = ""
_last_tool          = ""
_tool_count         = 0
_last_tool_ts       = 0.0
_progress_on_screen = False   # True while an inline progress line is showing
_submit_log_pos     = 0       # byte offset into submit_progress.log already displayed
_search_log_pos     = 0       # byte offset into search_progress.log already displayed


def _clear_progress() -> None:
    global _last_tool, _tool_count, _last_tool_ts
    _last_tool = ""
    _tool_count = 0
    _last_tool_ts = 0.0


def _poll_submit_log() -> None:
    """Print any new lines written to submit_progress.log since last poll."""
    global _submit_log_pos
    resume_output_dir = os.environ.get("RESUME_OUTPUT_DIR", str(Path.home() / "Documents" / "jobs" / "tailored")).rstrip("/")
    log_path = Path(resume_output_dir) / "state" / "submit_progress.log"
    try:
        if not log_path.exists():
            return
        with open(log_path, "r") as f:
            f.seek(_submit_log_pos)
            new = f.read()
            _submit_log_pos = f.tell()
        for line in new.splitlines():
            if line.strip():
                _flush_line(line)
    except OSError:
        pass


def _poll_search_log() -> None:
    """Print any new lines written to search_progress.log since last poll."""
    global _search_log_pos
    log_path = WORKSPACE / "state" / "search_progress.log"
    try:
        if not log_path.exists():
            return
        with open(log_path, "r") as f:
            f.seek(_search_log_pos)
            new = f.read()
            _search_log_pos = f.tell()
        for line in new.splitlines():
            if line.strip():
                _flush_line(line)
    except OSError:
        pass


def _show_progress(tool: str) -> None:
    import datetime, time
    global _last_tool, _tool_count, _last_tool_ts, _progress_on_screen
    now = time.time()
    _tool_count += 1
    if tool == _last_tool and _tool_count < 10 and (now - _last_tool_ts) < 5:
        return
    _last_tool    = tool
    _last_tool_ts = now
    ts = datetime.datetime.now().strftime('%H:%M:%S')
    count_str = f" ×{_tool_count}" if _tool_count > 1 else ""
    # Print inline (no newline) so real content can overwrite it
    print(f"{_ERASE_LINE}{_DIM}{ts} [{tool}{count_str}]{_RESET}", end="", flush=True)
    _progress_on_screen = True
    _tool_count = 0


def _flush_line(line: str) -> None:
    import datetime
    global _progress_on_screen
    _clear_progress()
    # If a progress line is showing inline, erase it before printing real content
    prefix = _ERASE_LINE if _progress_on_screen else ""
    _progress_on_screen = False
    clean = _ANSI_RE.sub('', line)
    if not clean.strip():
        print(prefix)
        return
    for tag, icon in _ICONS.items():
        clean = clean.replace(tag, icon)
    ts = f"{_DIM}{datetime.datetime.now().strftime('%H:%M:%S')}{_RESET} "
    if re.search(r'\*\*Progress:', clean):
        print(f"{prefix}{ts}{_MAGENTA}{clean}{_RESET}")
        return
    if re.match(r'\[task:', clean):
        print(f"{prefix}{ts}{_DIM}{clean}{_RESET}")
        return
    colored = clean.replace('✓', f'{_GREEN}✓{_RESET}')
    colored = colored.replace('✗', f'{_RED}✗{_RESET}')
    print(f"{prefix}{ts}{colored}")

def _stream_text(text: str) -> None:
    global _line_buf
    _line_buf += text
    while "\n" in _line_buf:
        line, _line_buf = _line_buf.split("\n", 1)
        _flush_line(line)

from .base import base_env, WORKSPACE, SKILLS_DIR, REPO_ROOT
from .search import definition as search_def
from .tailor import definition as tailor_def
from .submitter import definition as submitter_def

_DEFAULT_LOGIN_TIMEOUT = 60   # seconds


def _tracker_get_setting(key: str) -> str | None:
    """Read a value from the ## Settings table in the tracker."""
    tracker = WORKSPACE / "job_application_tracker.md"
    if not tracker.exists():
        return None
    import re
    text = tracker.read_text()
    m = re.search(rf'^\|\s*{re.escape(key)}\s*\|\s*([^|]+?)\s*\|', text, re.MULTILINE)
    return m.group(1).strip() if m else None


def _tracker_set_setting(key: str, value: str) -> None:
    """Write/update a value in the ## Settings table in the tracker, creating the section if needed."""
    import re
    tracker = WORKSPACE / "job_application_tracker.md"
    if not tracker.exists():
        return
    text = tracker.read_text()
    row = f"| {key} | {value} |"
    # Update existing row
    pattern = rf'^\|\s*{re.escape(key)}\s*\|[^|]*\|'
    if re.search(pattern, text, re.MULTILINE):
        text = re.sub(pattern, row, text, flags=re.MULTILINE)
    elif "## Settings" in text:
        # Append row to existing Settings table
        text = re.sub(
            r'(## Settings.*?\|[-| ]+\|)(.*?)(\n## |\Z)',
            lambda m: m.group(1) + m.group(2) + f"\n{row}" + m.group(3),
            text, flags=re.DOTALL
        )
    else:
        # Add Settings section before ## Jobs
        settings_block = (
            "\n## Settings\n\n"
            "| Key | Value |\n"
            "|-----|-------|\n"
            f"{row}\n"
        )
        text = re.sub(r'(\n## Jobs)', settings_block + r'\1', text)
    tracker.write_text(text)



SYSTEM_PROMPT = f"""\
You are the Job Autopilot orchestrator. You coordinate a three-stage pipeline
by delegating to specialized subagents using the Agent tool.

## Pipeline stages

1. **job-search** — Search job boards, filter results, write to tracker.
   Invoke when: user wants to find new jobs.
   After the search agent returns, relay its output line by line using `[Search]` prefix:
     [Search] ✓ Goldman Sachs — Quant Developer  (shortlisted; ...)
     [Search] ✗ JPMorgan — Junior Analyst  (screen_reject: too junior)
   Then print a one-line summary: jobs found, shortlisted, rejected.

2. **resume-tailor** — Writes tailored .md files (resume + cover letter) for each shortlisted job.
   The tailor agent produces ONLY .md files. After each tailor call returns, YOU must:

   a) Convert resume .md → .docx:
      python3 $MD_TO_DOCX_SCRIPT "<resume.md>" "$RESUME_TEMPLATE" "<resume.docx>"
      Then update the tracker: set resume_path = "<resume.docx>" (replace .md extension with .docx)

   b) Verify tables > 0:
      python3 -c "from docx import Document; d=Document('<resume.docx>'); assert len(d.tables)>0, 'BAD FORMAT'"

   b2) Convert cover letter .md → .docx (plain text, no template):
      python3 - << 'PYEOF'
      import sys
      from pathlib import Path
      from docx import Document
      from docx.shared import Pt, Inches
      md = Path("<cover_letter.md>").read_text()
      doc = Document()
      for sec in doc.sections:
          sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
          sec.left_margin = Inches(1); sec.right_margin = Inches(1)
      style = doc.styles['Normal']
      style.font.name = 'Calibri'; style.font.size = Pt(11)
      for line in md.strip().split('\n'):
          doc.add_paragraph(line)
      doc.save("<cover_letter.docx>")
      PYEOF
      Then update the tracker: set cover_letter_path = "<cover_letter.docx>" (replace .md extension with .docx)

   c) Check resume .md content against the source resume in $RESUME_DIR:
      - Line 1 name must exactly match the name in the source resume
      - All companies and roles listed must exist in the source resume (no invented employers)
      - EARLIER EXPERIENCE section must include all employers beyond the top 5 jobs
      - EARLIER EXPERIENCE entries must be ONE LINE EACH — no bullets, no second line (grep -A1 "EARLIER EXPERIENCE" to check)
      - Bullet point facts (tools, metrics, product names) must be traceable to the source resume
      - No invented details in any bullet

   d) If any check fails: ask tailor to fix the .md, then re-convert.

   e) Print one result line per job (see output format below).

   Invoke tailor in batches of up to 10 jobs. After each batch returns, run steps a–e for
   every job in the batch, then launch the next batch. While waiting for tailor, submitter
   and search can be running in parallel.

3. **job-submitter** — Submits applications via iterative code-act loop (write actions JSON →
   run the browser runner → read page state → repeat). Handles its own execution loop internally.
   Invoke when: tracker has `resume_ready` or `blocked` entries.
   After it returns, relay its output line by line using `[Submitter]` prefix.

## Shortlist target

`JOB_SEARCH_SHORTLIST_TARGET` is the max number of active jobs in the pipeline at once.
Before invoking job-search, count rows with status `shortlist` + `tailoring` + `resume_ready`.
If that total is already ≥ target, skip search entirely — there is already enough work queued.

If the user's prompt contains a shortlist count (e.g. "stop after 10 shortlists", "find 5 jobs",
"shortlist-target 20"), extract that number and override the env var for this run only.

## Startup — open blocked job pages for manual handling

Immediately after reading the tracker (step 1 below), before launching any agents:

1. Find the first 5 `blocked` entries in the tracker that have a URL in their Notes field.
2. For each blocked job, find the actual stuck URL by grepping the progress log:
   ```bash
   grep "<job_id>" "$RESUME_OUTPUT_DIR/state/submit_progress.log" | grep "ended on" | tail -1
   ```
   The line contains the URL after the colon (e.g. `ended on captcha: https://...`).
   Fall back to `$RESUME_OUTPUT_DIR/state/<job_id>/tab_url.txt`, then the Notes URL.
3. Open them in the background using Bash:
   ```bash
   ~/voracle-env/bin/python3 "$BROWSER_LOGIN_SCRIPT" \
     "$SUBMIT_BROWSER_PROFILE_DIR" \
     '["url1","url2","url3","url4","url5"]' &
   ```
4. Print one line per URL opened:
   `[Orchestrator] 🌐 Opened for manual review: Company — Role (url)`
5. Continue immediately — do not wait for the user.

This lets the user log in and submit blocked jobs manually while the pipeline runs autonomously.

## Your responsibilities — pipeline parallelism

**Do NOT run stages sequentially.** Run all stages that have work simultaneously by calling
multiple Agent tools in the same turn. The pipeline is a conveyor belt, not a waterfall:

1. Read the tracker once at startup.
2. In a SINGLE turn, call Agent tools for every stage that has work:
   - `job-search`    if shortlist count < target
   - `resume-tailor` if tracker has any `shortlist` entries
   - `job-submitter` if tracker has any `resume_ready` or `blocked` entries
   After submitter returns, execute the generated scripts via Bash (sequential, one at a time).
3. After each agent call returns, process its output (docx conversion, relay lines, etc.),
   then immediately re-launch that agent if there is still work for it.
4. Keep all three agents running concurrently whenever possible.
5. Loop until job-search finds zero new jobs AND tailor queue is empty AND submit queue is empty.

**Never wait for search to finish before starting tailor.**
**Never wait for tailor to finish before starting submitter.**
The moment the tracker has `shortlist` entries → tailor should be running.
The moment the tracker has `resume_ready` entries → submitter should have its browser open.

**NEVER ask for confirmation before any stage — run fully autonomously.**
`blocked` entries are treated exactly like `resume_ready` — launch submitter immediately, no questions asked.
If a stage fails or produces no results, skip it and continue.
Report a brief summary after each stage completes.

## Output format

Prefix every line of output with who is speaking:

  [Orchestrator] Starting tailor stage — 85 jobs
  [Tailor] ✓ Acme Corp — Quant Developer
  [Tailor] ✗ Firm B — Risk VP  (name on line 1 didn't match source, fixed and retried)
  [Orchestrator] Tailor stage complete — 84 done, 1 error
  [Orchestrator] Converting Acme_Corp_Quant_Developer_Resume_2026.md → .docx ... ok (5 tables)
  [Orchestrator] Content check passed: name ok, all employers found, no fabrications

Rules:
- [Orchestrator] for your own planning, conversion steps, checks, and stage summaries
- [Tailor] for each result line reported back from the tailor subagent
- [Submitter] for each result line reported back from the submitter subagent
- [Search] for each result line reported back from the search subagent
- One line per job — no multi-line blocks between jobs
- Stage icons: 🎬 Orchestrator · 🔍 Search · ✂️ Tailor · 📨 Submitter — use these everywhere, never 📄 for tailor

## Read-only browser tasks — reuse the submitter's authenticated browser

The submitter's Playwright browser is a persistent, already-signed-in session for whatever the user regularly uses: Gmail, Outlook, LinkedIn, employer portals, ATS dashboards, webmail, Calendar, etc. It accumulates cookies across runs, so by the time you need to read something, the user is almost certainly already logged in there.

If the user asks you to **look something up** rather than apply to jobs — check inbox, read a LinkedIn message thread, look at an interview calendar invite, verify an application status page, scrape an offer letter, open a webmail provider, check a recruiter reply, etc. — delegate to **job-submitter** with a read-only prompt like:

> "Read-only task — do NOT submit any application. Browse the site using the browser runner. Report findings back to me."

**Email scanning — read the BODY, not just the subject line.** Subject lines and snippets are not enough. Job-related emails often contain critical details inside the body: OA deadlines, interview scheduling links, background check forms, offer details, next-step instructions. When checking email:

1. **Round 1 — list view:** Navigate to the inbox/search URL, wait for load, `evaluate` to extract a list of rows with sender, subject, snippet, and a clickable selector (e.g. Gmail `.zA` row index or element ID).
2. **Round 2+ — open each relevant email:** For every email that looks job-related from the list view, click into it (one per round or batch if possible), `evaluate` to extract the full visible body text, then navigate back to the list. Look for:
   - Action items: OA links, scheduling links, "complete by" deadlines, forms to fill
   - Interview details: date/time, interviewer names, preparation instructions
   - Offer/rejection signals: "we'd like to extend", "unfortunately", "next steps"
   - Status updates: "your application has moved to", "background check initiated"
3. **Report** with a clear split: ACTION NEEDED (with deadlines) vs. informational (confirmations, rejections, surveys).

Do NOT skip opening emails — a subject line like "Your application update" could be a rejection, an interview invite, or an OA. You can't tell without reading the body.

**After the submitter finishes scanning, update the tracker yourself.** Read `$TRACKER_PATH`, match each email finding to the corresponding tracker row by company + role, and update:

| Email signal | Tracker update |
|---|---|
| Rejection email ("unfortunately", "not moving forward", "after careful consideration") | Status → `denied`. Append `DENIED <date>: rejection email received.` to notes. (`denied` = employer rejected; `screen_reject`/`user_reject` = we filtered it ourselves) |
| OA / coding assessment invite | Keep status `applied`. Append `OA RECEIVED <date>: <platform> assessment, deadline <date if found>.` to notes. |
| Interview invite / scheduling link | Status → `interviewing`. Append `INTERVIEW <date>: <details>.` to notes. |
| Offer | Status → `offer`. Append `OFFER <date>: <details>.` to notes. |
| "Application received" confirmation (for a row still in `blocked` or `resume_ready`) | Status → `applied`. Append `Confirmed applied via email <date>.` to notes. |
| Draft / incomplete application reminder | Keep current status. Append `DRAFT REMINDER <date>: application saved but not submitted.` to notes. |

Match against **every row in the tracker** — do NOT restrict to a subset or "target companies". Every company the user has applied to is in the tracker and deserves an update if there is a matching email.

"Conservatively" means: only update a specific row when you are confident the email matches that company + role. If a rejection email says "Senior Lead AI Engineer" and there are 3 Capital One AI roles, update only the one(s) whose title matches. When unsure which row, note the ambiguity in all candidate rows and let the user decide. It does NOT mean skip companies.

Build the URL and evaluate queries from the user's request. Examples:
- Gmail search: `https://mail.google.com/mail/u/0/#search/<query>` + evaluate `.zA` rows, then click into each
- Outlook: `https://outlook.live.com/mail/0/` + evaluate message list, then click into each
- LinkedIn messages: `https://www.linkedin.com/messaging/` + click into conversation cards
- Any employer ATS status dashboard the user already logged into

The submitter agent knows how to drive Playwright and how to parse `page.interactive` / `evaluate` results. Its browser state directory persists across runs.

**Do NOT ask the user to run `/mcp`** or authenticate any MCP server (Gmail, Outlook, Calendar, etc.) for read-only lookup tasks. **Do NOT use `mcp__*` tools for these tasks.** The Playwright route via submitter is the only one that works in this pipeline — it reuses existing logins and doesn't require fresh OAuth every time.

MCP tools may still be appropriate for things that genuinely need a programmatic API (e.g. bulk structured calendar writes), but default to the browser route for anything read-only.

## User overrides — listen to the user

If the user's prompt explicitly limits scope (e.g. "search only", "do not run submitter",
"skip tailor", "only tailor today"), HONOR it for this run. Skip the disabled stages
even if the tracker has work for them. The default parallel-pipeline behavior applies
only when the user gives a generic prompt with no scope restriction.

## Critical rules — do NOT violate

- **Never check the filesystem yourself to decide whether tailoring is needed.** If the tracker has `shortlist` entries, invoke the resume-tailor subagent — always. Do not inspect `$RESUME_OUTPUT_DIR` or determine that files "already exist". That is the tailor agent's job.
- **Never mark jobs `resume_ready` yourself.** Only the tailor subagent may do this.
- **Never mark jobs `blocked` or `applied` yourself** — EXCEPT when updating based on email scan results (rejections, interview invites, OA notifications, confirmed-applied signals). In that case the orchestrator updates the tracker directly after the submitter reports email findings.

## Tracker
{WORKSPACE / "job_application_tracker.md"}
"""


# Fallback chain — when a model hits its rate limit, advance to the next.
# Both short aliases (passed to subagents via AgentDefinition) and full IDs
# (used for the orchestrator's own model) are mapped.
_MODEL_FALLBACK = {
    "sonnet":             "opus",
    "claude-sonnet-4-6":  "claude-opus-4-7",
    "opus":               "haiku",
    "claude-opus-4-7":    "claude-haiku-4-5-20251001",
    "haiku":              None,
    "claude-haiku-4-5-20251001": None,
}


def _is_rate_limit(text: str) -> bool:
    """Detect a Claude Code subscription rate-limit message in streamed text."""
    t = text.lower()
    return ("limit" in t and "resets" in t) or "hit your limit" in t


def _rate_limit_wait_secs(err: str) -> int:
    """Parse 'resets H[am/pm] (Timezone)' from a rate-limit error and return seconds to wait."""
    m = re.search(r'resets\s+(\d+)([ap]m)\s+\(([^)]+)\)', err, re.I)
    if m:
        hour = int(m.group(1))
        if m.group(2).lower() == 'pm' and hour != 12:
            hour += 12
        elif m.group(2).lower() == 'am' and hour == 12:
            hour = 0
        try:
            tz = ZoneInfo(m.group(3))
            now = datetime.datetime.now(tz)
            reset = now.replace(hour=hour, minute=5, second=0, microsecond=0)
            if reset <= now:
                reset += datetime.timedelta(days=1)
            return max(60, int((reset - now).total_seconds()))
        except Exception:
            pass
    return 30 * 60  # default: wait 30 minutes


async def run(prompt: str, stream: bool = True, headed: bool = True) -> str:
    env = base_env()
    env.update({
        "JOB_SEARCH_KEYWORDS":    os.environ.get("JOB_SEARCH_KEYWORDS", ""),
        "JOB_SEARCH_LOCATION":    os.environ.get("JOB_SEARCH_LOCATION", ""),
        "JOB_SEARCH_HANDOFF":     os.environ.get("JOB_SEARCH_HANDOFF", str(WORKSPACE / "SEARCH_AGENT_HANDOFF.md")),
        "JOB_SEARCH_MIN_SALARY":  os.environ.get("JOB_SEARCH_MIN_SALARY", ""),
        "JOB_SEARCH_MAX_AGE_DAYS": os.environ.get("JOB_SEARCH_MAX_AGE_DAYS", "90"),
        "JOB_SEARCH_SHORTLIST_TARGET": os.environ.get("JOB_SEARCH_SHORTLIST_TARGET", "30"),
        "USER_FIRST_NAME":  os.environ.get("USER_FIRST_NAME", ""),
        "USER_LAST_NAME":   os.environ.get("USER_LAST_NAME", ""),
        "USER_EMAIL":       os.environ.get("USER_EMAIL", ""),
        "USER_PHONE":       os.environ.get("USER_PHONE", ""),
        "USER_LINKEDIN":    os.environ.get("USER_LINKEDIN", ""),
        "RESUME_OUTPUT_DIR": os.environ.get("RESUME_OUTPUT_DIR", str(Path.home() / "Documents" / "jobs" / "tailored") + "/"),
        "RESUME_TEMPLATE":   os.environ.get("RESUME_TEMPLATE", str(SKILLS_DIR / "tailor" / "scripts" / "sample_placeholders.docx")),
        "MD_TO_DOCX_SCRIPT": os.environ.get("MD_TO_DOCX_SCRIPT", str(SKILLS_DIR / "tailor" / "scripts" / "md_to_docx.py")),
        "TRACKER_PATH":     os.environ.get("TRACKER_PATH", str(WORKSPACE / "job_application_tracker.md")),
        "SKILLS_DIR":       str(SKILLS_DIR),
        "SUBMIT_RUNNER":    str(SKILLS_DIR / "submitter" / "scripts" / "submit_runner.py"),
        "CREDENTIALS_FILE": os.environ.get("CREDENTIALS_FILE", str(WORKSPACE / "credentials.md")),
        "SEARCH_PROGRESS_LOG": str(WORKSPACE / "state" / "search_progress.log"),
        "BROWSER_LOGIN_SCRIPT":      str(REPO_ROOT / "scripts" / "browser_login.py"),
        "BROWSER_RESTART_SCRIPT":    str(REPO_ROOT / "scripts" / "browser_restart.py"),
        "SUBMIT_BROWSER_PROFILE_DIR": str(Path.home() / ".jobautopilot" / "browser_profiles" / "submit"),
        "USER_GENDER":      os.environ.get("USER_GENDER", ""),
        "USER_RACE":        os.environ.get("USER_RACE", ""),
        "USER_HISPANIC":    os.environ.get("USER_HISPANIC", ""),
        "USER_VETERAN":     os.environ.get("USER_VETERAN", ""),
        "USER_DISABILITY":  os.environ.get("USER_DISABILITY", ""),
        "USER_WORK_AUTH":   os.environ.get("USER_WORK_AUTH", ""),
        "USER_NEED_SPONSOR": os.environ.get("USER_NEED_SPONSOR", ""),
        "USER_NON_COMPETE": os.environ.get("USER_NON_COMPETE", ""),
    })

    # Resolve login timeout: CLI override > tracker > default
    if os.environ.get("LOGIN_HUMAN_TIMEOUT_OVERRIDE"):
        login_timeout = int(os.environ["LOGIN_HUMAN_TIMEOUT"])
        _tracker_set_setting("login_timeout", str(login_timeout))
    else:
        saved = _tracker_get_setting("login_timeout")
        login_timeout = int(saved) if saved and saved.isdigit() else _DEFAULT_LOGIN_TIMEOUT
        os.environ.setdefault("LOGIN_HUMAN_TIMEOUT", str(login_timeout))
    env["LOGIN_HUMAN_TIMEOUT"] = str(login_timeout)

    # Seek to end of existing progress logs so we only show NEW lines from this run
    global _submit_log_pos, _search_log_pos
    resume_output_dir = os.environ.get("RESUME_OUTPUT_DIR", str(Path.home() / "Documents" / "jobs" / "tailored")).rstrip("/")
    for _log_path, _pos_var in [
        (Path(resume_output_dir) / "state" / "submit_progress.log", "_submit_log_pos"),
        (WORKSPACE / "state" / "search_progress.log",               "_search_log_pos"),
    ]:
        try:
            if _log_path.exists():
                globals()[_pos_var] = _log_path.stat().st_size
        except OSError:
            pass
    (WORKSPACE / "state").mkdir(parents=True, exist_ok=True)

    _flush_line(f"[Orchestrator] starting — may take a minute to load, please be patient...")

    result = ""
    # Track current model for orchestrator + each subagent; advanced down
    # _MODEL_FALLBACK on rate-limit hits.
    orch_model      = "claude-sonnet-4-6"
    search_model    = "haiku"
    tailor_model    = "opus"
    submitter_model = "sonnet"

    while True:
        # Captured rate-limit text from streamed output. The SDK raises a
        # generic "Command failed with exit code 1" exception on rate limit,
        # so we sniff the streamed TextBlock to know what actually happened.
        rate_limit_text = ""
        try:
            async for message in query(
                prompt=prompt,
                options=ClaudeAgentOptions(
                    cwd=str(WORKSPACE),
                    allowed_tools=["Read", "Bash", "Agent"],
                    system_prompt=SYSTEM_PROMPT,
                    permission_mode="bypassPermissions",
                    model=orch_model,
                    env=env,
                    agents={
                        "job-search":    search_def(headed, model=search_model),
                        "resume-tailor": tailor_def(headed, model=tailor_model),
                        "job-submitter": submitter_def(True, model=submitter_model),  # always headed — user may need to solve CAPTCHAs
                    },
                ),
            ):
                if isinstance(message, AssistantMessage) and stream:
                    for block in message.content:
                        if isinstance(block, TextBlock):
                            if _is_rate_limit(block.text):
                                rate_limit_text = block.text
                            _stream_text(block.text)
                        elif isinstance(block, ThinkingBlock):
                            pass  # internal reasoning — not shown
                        elif isinstance(block, ToolUseBlock):
                            if block.name == "Agent":
                                agent_name = block.input.get("name", "subagent")
                                _stream_text(f"[Orchestrator] → launching {agent_name}...\n")
                            elif block.name in ("Bash", "Read"):
                                pass  # too noisy to print every read/bash
                elif isinstance(message, TaskStartedMessage) and stream:
                    _clear_progress()
                    _stream_text(f"[task: {message.description}]\n")
                elif isinstance(message, TaskProgressMessage) and stream:
                    _poll_submit_log()
                    _poll_search_log()
                    if message.last_tool_name:
                        _show_progress(message.last_tool_name)
                elif isinstance(message, ResultMessage):
                    result = message.result
                    if stream:
                        if _line_buf:
                            _flush_line(_line_buf)
                        print()
            break  # completed without rate-limit error
        except Exception as e:
            err = str(e)
            hit = rate_limit_text if rate_limit_text else (err if _is_rate_limit(err) else "")
            if hit:
                # Try advancing every model that has a fallback. Switching all
                # of them is the simplest correct behavior — we don't know
                # exactly which model triggered the limit (the streamed text
                # doesn't always say), and shared subscription tiers like
                # "sonnet" are usually exhausted in lockstep across agents.
                advanced = False
                if _MODEL_FALLBACK.get(orch_model):
                    orch_model = _MODEL_FALLBACK[orch_model]
                    advanced = True
                if _MODEL_FALLBACK.get(search_model):
                    search_model = _MODEL_FALLBACK[search_model]
                    advanced = True
                if _MODEL_FALLBACK.get(tailor_model):
                    tailor_model = _MODEL_FALLBACK[tailor_model]
                    advanced = True
                if _MODEL_FALLBACK.get(submitter_model):
                    submitter_model = _MODEL_FALLBACK[submitter_model]
                    advanced = True
                if advanced:
                    _flush_line(
                        f"[Orchestrator] Rate limit — falling back to "
                        f"orch={orch_model}, search={search_model}, "
                        f"tailor={tailor_model}, submitter={submitter_model}"
                    )
                    continue
                # No more fallbacks — sleep until the limit resets, then retry.
                secs = _rate_limit_wait_secs(hit)
                _flush_line(f"[Orchestrator] Rate limit — all fallbacks exhausted, resuming in {secs // 60} min...")
                await anyio.sleep(secs)
                continue
            raise

    return result
